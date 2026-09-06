"""文档加载器：支持 TXT / MD / PDF / DOCX，统一返回 LangChain Document 列表。"""
from __future__ import annotations

import os
from typing import List

from langchain_core.documents import Document
from langchain_community.document_loaders import PyPDFLoader, TextLoader


SUPPORTED_EXTS = {".txt", ".md", ".markdown", ".pdf", ".docx"}


class UnsupportedFormatError(ValueError):
    pass


def _load_docx(path: str, source: str) -> List[Document]:
    from docx import Document as DocxDocument

    doc = DocxDocument(path)
    parts: List[str] = []
    # 遍历段落与表格，尽量保留结构信息
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    content = "\n".join(parts)
    return [Document(page_content=content, metadata={"source": source})]


def _load_doc(path: str) -> List[Document]:
    ext = os.path.splitext(path)[1].lower()
    source = os.path.basename(path)
    if ext in (".txt",):
        loader = TextLoader(path, encoding="utf-8")
        docs = loader.load()
        for d in docs:
            d.metadata["source"] = source
        return docs
    if ext in (".md", ".markdown"):
        with open(path, "r", encoding="utf-8") as f:
            content = f.read()
        return [Document(page_content=content, metadata={"source": source})]
    if ext == ".pdf":
        loader = PyPDFLoader(path)
        return loader.load()
    if ext == ".docx":
        return _load_docx(path, source)
    raise UnsupportedFormatError(f"不支持的文件类型: {ext}，请使用 {sorted(SUPPORTED_EXTS)}")


def load_document(path: str) -> List[Document]:
    """加载单个文档。找不到文件或格式不支持时抛出明确异常。"""
    if not os.path.exists(path):
        raise FileNotFoundError(f"文件不存在: {path}")
    if os.path.splitext(path)[1].lower() not in SUPPORTED_EXTS:
        raise UnsupportedFormatError(
            f"不支持的文件类型，请使用 {sorted(SUPPORTED_EXTS)}"
        )
    return _load_doc(path)


def load_directory(dir_path: str) -> List[Document]:
    """递归加载目录下所有支持的文档。"""
    docs: List[Document] = []
    for root, _, files in os.walk(dir_path):
        for name in files:
            if os.path.splitext(name)[1].lower() in SUPPORTED_EXTS:
                fp = os.path.join(root, name)
                try:
                    docs.extend(_load_doc(fp))
                except Exception as e:  # noqa: BLE001 单个文件失败不阻塞整体
                    print(f"[load_directory] 跳过 {name}: {e}")
    return docs